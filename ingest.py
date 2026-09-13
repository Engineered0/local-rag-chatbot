import os, json, time, hashlib, chromadb, pdfplumber
from config import CHUNK_SIZE
from describe import ocr, describe, clear_cache
from embedder import embed

BASE = os.path.dirname(os.path.abspath(__file__))
DOCS = f"{BASE}/docs"
MANIFEST = f"{BASE}/.ingest_cache.json"

TEXT_EXT = {".txt", ".md", ".csv", ".json", ".py", ".log", ".rst"}
IMG_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}

col = chromadb.PersistentClient(f"{BASE}/chroma_db").get_or_create_collection("docs")


# ----------------------------------------------------------------- helpers

def is_junk(text):
    if len(text.split()) < 20:
        return True
    return text.count(".") > len(text) * 0.15


def chunk_text(text, size=CHUNK_SIZE):
    paras = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks, cur = [], []
    for p in paras:
        cur.append(p)
        if sum(len(x.split()) for x in cur) >= size:
            chunks.append("\n\n".join(cur))
            cur = cur[-1:]
    if cur:
        chunks.append("\n\n".join(cur))
    return chunks


def flatten_table(table, name, page):
    out = []
    header = [str(c or "").strip() for c in table[0]]
    for row in table[1:]:
        cells = [str(c or "").strip() for c in row]
        pairs = ", ".join(f"{h}: {v}" for h, v in zip(header, cells) if h and v)
        if pairs:
            out.append(f"[{name} p{page}] {pairs}")
    return out


def file_hash(path):
    h = hashlib.md5()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(65536), b""):
            h.update(block)
    return h.hexdigest()


# ----------------------------------------------------------------- loaders

def process_pdf(path, name):
    records = []
    t_start = time.time()
    with pdfplumber.open(path) as pdf:
        total = len(pdf.pages)
        for pno, page in enumerate(pdf.pages, 1):
            t0 = time.time()
            kinds = []

            for table in page.extract_tables():
                for line in flatten_table(table, name, pno):
                    records.append((line, pno, "table"))
                    kinds.append("table")

            text = (page.extract_text() or "").strip()
            has_figure = len(page.images) > 0

            if text:
                for c in chunk_text(text):
                    if not is_junk(c):
                        records.append((f"[{name} p{pno}] {c}", pno, "text"))
                        kinds.append("text")

            ocr_text = ""
            if has_figure or not text:
                ocr_text = ocr(path, pno)
                if len(ocr_text.split()) > 15:
                    kind = "figure_ocr" if text else "ocr"
                    records.append((f"[{name} p{pno}] Figure text: {ocr_text}",
                                    pno, kind))
                    kinds.append(kind)

            if has_figure and len(ocr_text.split()) < 15:
                desc = describe(path, pno)
                if desc:
                    records.append((f"[{name} p{pno}] Figure description: {desc}",
                                    pno, "vision"))
                    kinds.append("vision")
                elif not text:
                    records.append((
                        f"[{name} p{pno}] Image-only page. Open the PDF here.",
                        pno, "image_only"))
                    kinds.append("image_only")

            clear_cache()

            dt = time.time() - t0
            elapsed = time.time() - t_start
            eta = (elapsed / pno) * (total - pno)
            tag = "+".join(sorted(set(kinds))) or "empty"
            print(f"  p{pno}/{total}  {dt:5.1f}s  {tag:<24} "
                  f"eta {eta/60:4.1f}m", end="\r", flush=True)

    print(" " * 78, end="\r")
    return records


def process_image(path, name):
    records = []

    t = ocr(path)
    if len(t.split()) > 10:
        records.append((f"[{name}] Text in image: {t}", 0, "ocr"))

    desc = describe(path)
    if desc:
        records.append((f"[{name}] Image description: {desc}", 0, "vision"))

    if not records:
        records.append((f"[{name}] Image with no readable content.",
                        0, "image_only"))

    clear_cache()
    return records


def process_docx(path, name):
    from docx import Document
    d = Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n\n".join(parts)
    return [(f"[{name}] {c}", 0, "text")
            for c in chunk_text(text) if not is_junk(c)]


def load(path, name, ext):
    if ext == ".pdf":
        return process_pdf(path, name)
    if ext in IMG_EXT:
        return process_image(path, name)
    if ext == ".docx":
        return process_docx(path, name)
    if ext in TEXT_EXT:
        text = open(path, encoding="utf-8", errors="replace").read()
        return [(f"[{name}] {c}", 0, "text")
                for c in chunk_text(text) if not is_junk(c)]
    return None


# -------------------------------------------------------------------- main

if __name__ == "__main__":
    t_run = time.time()
    manifest = json.load(open(MANIFEST)) if os.path.exists(MANIFEST) else {}
    present = set()

    for name in sorted(os.listdir(DOCS)):
        if name.startswith("."):
            continue
        path = f"{DOCS}/{name}"
        ext = os.path.splitext(name)[1].lower()
        present.add(name)

        h = file_hash(path)
        if manifest.get(name) == h:
            print(f"{name}: unchanged, skipped")
            continue

        print(f"{name}: processing...")
        t_file = time.time()
        try:
            records = load(path, name, ext)
        except Exception as e:
            print(f"  FAILED: {e}")
            continue
        t_extract = time.time() - t_file

        if records is None:
            print(f"  unsupported type ({ext})")
            continue
        if not records:
            print("  nothing extracted")
            continue

        texts = [r[0] for r in records]
        t_embed_start = time.time()
        vectors = []
        for i in range(0, len(texts), 32):
            vectors.extend(embed(texts[i:i + 32]))
            done = min(i + 32, len(texts))
            rate = done / max(time.time() - t_embed_start, 0.01)
            print(f"  embedding {done}/{len(texts)}  {rate:.0f}/s   ",
                  end="\r", flush=True)
        t_embed = time.time() - t_embed_start

        # clear this file's old chunks so stale kinds don't linger
        col.delete(where={"source": name})
        col.upsert(
            ids=[f"{name}-p{r[1]}-{r[2]}-{i}" for i, r in enumerate(records)],
            documents=texts,
            embeddings=vectors,
            metadatas=[{"source": name, "page": r[1], "kind": r[2]}
                       for r in records],
        )

        manifest[name] = h
        counts = {}
        for _, _, k in records:
            counts[k] = counts.get(k, 0) + 1
        print(f"  {len(records)} chunks {counts}        ")
        print(f"  extract {t_extract:.1f}s · embed {t_embed:.1f}s · "
              f"total {time.time() - t_file:.1f}s")

    for gone in set(manifest) - present:
        col.delete(where={"source": gone})
        del manifest[gone]
        print(f"{gone}: removed from database")

    json.dump(manifest, open(MANIFEST, "w"), indent=2)
    print(f"\ntotal: {col.count()} chunks in database")
    print(f"run took {(time.time() - t_run)/60:.1f} min")